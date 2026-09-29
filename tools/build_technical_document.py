"""Build the interview technical report as a polished DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path("deliverables/Allianz_Claims_Copilot_Technical_Approach.docx")
NAVY = "123047"
PALE_BLUE = "EAF2F7"
LIGHT_GRAY = "D9D9D9"
TEXT = RGBColor(30, 38, 45)


def set_cell_fill(cell: object, color: str) -> None:
    """Set a table cell background color."""

    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def set_cell_margins(
    cell: object,
    top: int = 100,
    start: int = 120,
    bottom: int = 100,
    end: int = 120,
) -> None:
    """Apply consistent cell padding in twips."""

    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table: object) -> None:
    """Add visible light-gray borders to a table."""

    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), LIGHT_GRAY)


def mark_header_row(row: object) -> None:
    """Mark a first row as a repeating semantic table header."""

    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def style_table(table: object, widths: list[float] | None = None) -> None:
    """Format a table for readable print output."""

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[column_index])
            if row_index == 0:
                set_cell_fill(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_fill(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else TEXT
                    run.bold = row_index == 0


def add_table(
    document: Document, headers: list[str], rows: list[list[str]], widths: list[float]
) -> None:
    """Add and style a comparison table."""

    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    style_table(table, widths)
    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    """Add a compact native bullet list."""

    for item in items:
        paragraph = document.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)


def add_page_number(paragraph: object) -> None:
    """Insert an editable PAGE field in a footer paragraph."""

    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])


def configure_document(document: Document) -> None:
    """Set the report page and type system."""

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.1

    title = document.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Heading 1", 19), ("Heading 2", 14)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    add_page_number(section.footer.paragraphs[0])


def build_document() -> Document:
    """Create the complete technical report."""

    document = Document()
    configure_document(document)

    document.add_heading("Agentic RAG for Motor Claims Handling", 0)
    subtitle = document.add_paragraph("Technical approach and five day implementation plan")
    subtitle.style = document.styles["Subtitle"]
    subtitle.runs[0].font.color.rgb = TEXT
    document.add_paragraph("AI Engineer technical assessment")
    document.add_paragraph("Source: Manual de criterios de las comisiones CIDE ASCIDE CICOS")
    document.add_paragraph("Prepared 29 September 2026")
    document.add_paragraph()
    opening = document.add_paragraph()
    opening.add_run("Main conclusion  ").bold = True
    opening.add_run(
        "A retrieval augmented system is the safest and most testable approach for this "
        "assessment. It can show the manual evidence behind each answer, update without "
        "model training and refuse conclusions that the retrieved pages do not support."
    )
    document.add_page_break()

    document.add_heading("Purpose and scope", level=1)
    document.add_paragraph(
        "I designed the solution for a claims handler who needs to ask factual questions "
        "or describe a motor accident. The assistant identifies the relevant parties and "
        "circumstances, explains responsibility under CIDE, ASCIDE or CICOS and cites the "
        "manual pages used. A person remains responsible for the final claim decision."
    )
    document.add_paragraph(
        "The system does not decide legal liability, insurance coverage, compensation, "
        "payment, rejection or criminal responsibility. This boundary matters because the "
        "manual defines criteria between insurers and the supplied edition dates from 2004."
    )

    document.add_heading("Assessment requirements", level=2)
    add_table(
        document,
        ["Requirement", "Implementation evidence"],
        [
            [
                "RAG over the supplied manual",
                "Page-aware ingestion, local embeddings and persistent Chroma search",
            ],
            ["LLM response generation", "Gemini on Vertex AI with a Pydantic output schema"],
            [
                "Accident analysis",
                "Parties, facts, agreement responsibility, missing information and citations",
            ],
            [
                "Basic quality evaluation",
                "Recall at K, MRR, citation validity and structured completeness",
            ],
            [
                "Source code and documentation",
                "API, UI, CLI, tests, Docker, CI, report and presentation",
            ],
        ],
        [2.2, 4.6],
    )

    document.add_heading("Architecture", level=1)
    document.add_paragraph(
        "The application separates ingestion, retrieval, orchestration, generation, "
        "validation and presentation. FastAPI, Streamlit and the command line call the same "
        "application service, so the business logic has one implementation."
    )
    add_table(
        document,
        ["Stage", "Responsibility", "Technology"],
        [
            [
                "Ingestion",
                "Extract, clean and split pages without crossing page boundaries",
                "pypdf and deterministic Python",
            ],
            ["Index", "Store text, vectors and source metadata", "Chroma with cosine HNSW"],
            [
                "Retrieve",
                "Return the closest six chunks and apply a distance threshold",
                "Multilingual MiniLM",
            ],
            [
                "Orchestrate",
                "Route, retrieve, generate, validate, retry once and stop",
                "LangGraph",
            ],
            [
                "Generate",
                "Return a typed Spanish response from retrieved evidence",
                "Gemini 2.5 Flash on Vertex AI",
            ],
            ["Serve", "Expose the same use case to API and demo UI", "FastAPI and Streamlit"],
        ],
        [1.0, 3.7, 2.1],
    )

    document.add_heading("Why RAG rather than fine tuning", level=2)
    document.add_paragraph(
        "The corpus contains one relatively small manual and the interview requires evidence "
        "that can be inspected live. Re-indexing handles a new manual version immediately. "
        "Fine-tuning would add cost and evaluation complexity while making source attribution "
        "harder. The LLM still contributes language understanding and structured synthesis."
    )

    document.add_heading("Ingestion and retrieval", level=1)
    document.add_paragraph(
        "The supplied PDF has 111 pages. The pipeline creates 134 chunks and detects 62 manual "
        "sections. Each identifier is derived from the source, page, position and cleaned text. "
        "Repeated ingestion therefore updates the same records instead of duplicating them."
    )
    document.add_paragraph(
        "The PDF renders correctly but its old embedded text replaces some accented characters "
        "with an unknown-glyph marker. The pipeline removes only that marker. It does not guess "
        "a replacement that could change a technical term. Page citations keep the rendered PDF "
        "as the source of truth. Spanish OCR is the first production improvement."
    )
    document.add_heading("Retrieval controls", level=2)
    add_bullets(
        document,
        [
            "Chunks never cross a PDF page, which makes every citation auditable.",
            "A multilingual model supports Spanish source text and natural-language questions.",
            "Top K and maximum distance are environment settings and can be "
            "tuned on the golden dataset.",
            "Retrieval evaluation runs without an LLM or external API call.",
        ],
    )

    document.add_heading("Agent workflow", level=1)
    document.add_paragraph(
        "LangGraph carries explicit state: query, deterministic route, retrieved chunks, typed "
        "answer, warnings, validation feedback and retry count. Each node has one responsibility."
    )
    add_table(
        document,
        ["Step", "Rule", "Failure behavior"],
        [
            ["Validate input", "Reject empty or very short requests", "Return a validation error"],
            [
                "Route",
                "Use business keywords for question or accident case",
                "Default to a factual question",
            ],
            [
                "Retrieve",
                "Keep results inside the relevance threshold",
                "Record that no evidence met the threshold",
            ],
            [
                "Generate",
                "Use only retrieved context and a structured schema",
                "Propagate transient provider errors",
            ],
            [
                "Validate",
                "Match every source, page and chunk identifier",
                "Return feedback to generation",
            ],
            ["Terminate", "Allow one retry", "Return a low-confidence human-review response"],
        ],
        [1.15, 3.45, 2.2],
    )

    document.add_heading("Grounding and safety", level=1)
    document.add_paragraph(
        "Retrieved text is marked as untrusted reference content. It cannot change the system "
        "policy or authorize an action. The LLM must separate agreement responsibility from "
        "legal liability and list facts that could change the answer. Code verifies the citation "
        "metadata after generation; fluent text alone does not pass validation."
    )
    add_bullets(
        document,
        [
            "No credentials or project identifiers are committed to Git.",
            "The MVP indexes only the supplied manual and does not require personal claim data.",
            "Retries cover only a citation-format failure and stop after one additional attempt.",
            "The UI displays the 2004 source limitation and requires specialist review.",
            "Logs in a production service must omit raw personal data and complete prompts.",
        ],
    )

    document.add_heading("Evaluation and testing", level=1)
    document.add_paragraph(
        "The versioned dataset contains ten questions. It covers the five interview scenarios "
        "and additional sections on semaphores, reverse motion, priority and direct collision. "
        "Relevant PDF pages provide labels for retrieval metrics."
    )
    add_table(
        document,
        ["Layer", "Metric or test", "Reason"],
        [
            [
                "Retrieval",
                "Recall at 6 and mean reciprocal rank",
                "Checks whether useful evidence is found and ranked early",
            ],
            [
                "Generation",
                "Citation validity and required-field completeness",
                "Detects unsupported references and broken output contracts",
            ],
            [
                "Logic",
                "Eleven unit tests with in-memory fakes",
                "Covers cleaning, chunking, routing, fallback and metrics without external APIs",
            ],
            [
                "Quality",
                "Ruff in local validation and CI",
                "Keeps implementation readable and consistent",
            ],
        ],
        [1.2, 2.7, 2.9],
    )
    document.add_paragraph(
        "Expert-labelled correctness and faithfulness should be added before operational use. "
        "Retrieval and generation remain separate measures because a well-written answer cannot "
        "compensate for missing or irrelevant evidence."
    )

    document.add_heading("Five day implementation plan", level=1)
    add_table(
        document,
        ["Day", "Deliverable", "Acceptance check"],
        [
            [
                "1",
                "Requirements, PDF inspection and architecture",
                "Pages extract with stable metadata and decisions are recorded",
            ],
            [
                "2",
                "Ingestion and retrieval",
                "Golden questions return relevant pages in the top six",
            ],
            ["3", "Agent graph and guardrails", "Citations validate and every route terminates"],
            ["4", "API, UI, Docker and evaluation", "Sample cases run through one shared service"],
            [
                "5",
                "Tests, deliverables and rehearsal",
                "Clean setup and a timed 30 to 45 minute presentation",
            ],
        ],
        [0.55, 2.85, 3.4],
    )

    document.add_heading("Risks and mitigations", level=1)
    add_table(
        document,
        ["Risk", "Impact", "Mitigation"],
        [
            [
                "Damaged PDF text layer",
                "Retrieval may miss accented terms",
                "Page citations now; Spanish OCR and checksum validation next",
            ],
            [
                "Manual dates from 2004",
                "Answer may be operationally outdated",
                "Visible warning, version metadata and assigned document owner",
            ],
            [
                "Invented citation",
                "Conclusion cannot be audited",
                "Validate citation metadata and fall back after one retry",
            ],
            [
                "Model or network failure",
                "Live demo interruption",
                "Pre-built index, health check and transparent fallback explanation",
            ],
            [
                "Five-day constraint",
                "Breadth may reduce quality",
                "One complete vertical slice and a documented backlog",
            ],
        ],
        [1.55, 2.15, 3.1],
    )

    document.add_heading("Production roadmap", level=1)
    document.add_paragraph(
        "The next release should add Spanish OCR, document effective dates, hybrid retrieval and "
        "reranking. A pilot would then introduce expert-labelled evaluations, access control, PII "
        "redaction, regional deployment, audit logs and traces for each graph node. Integration "
        "with a claims platform should stop before any consequential action until the business, "
        "legal, security and model-risk owners approve the control design."
    )

    document.add_heading("Repository evidence", level=1)
    add_bullets(
        document,
        [
            "src/allianz_rag contains ingestion, storage, graph, LLM, API, "
            "service and evaluation modules.",
            "evaluation/golden_dataset.jsonl contains the versioned cases and relevant pages.",
            "tests contains external-API-free unit tests for the important logic.",
            "Dockerfile, docker-compose.yml and .github/workflows/ci.yml define packaging and CI.",
            "docs/interview_runbook.md contains the timed demonstration and likely questions.",
        ],
    )

    return document


def main() -> None:
    """Write the final report."""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
